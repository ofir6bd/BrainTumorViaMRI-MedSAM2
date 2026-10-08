"""Builds Presentation/BrainTumour_YOLO_MedSAM2.pptx (15 slides) from the images in Presentation/images.

Needs python-pptx (pip install python-pptx). Run from the repo root:
    .venv\\Scripts\\python.exe Presentation\\make_charts.py
    .venv\\Scripts\\python.exe Presentation\\build_presentation.py
Every number on the slides comes from the saved run results (see outline.md for the sources).
"""
import os
import re

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "images")
OUT = os.path.join(HERE, "BrainTumour_YOLO_MedSAM2.pptx")

BG, PANEL, TEXT, MUTED = "0F1220", "1A1F33", "E8EAF2", "9AA0B8"
YOLO_C, MS_C, GOOD, LINE = "3987E5", "E06A2E", "2FB35C", "2A2F47"
FONT = "Segoe UI"
W, H = Inches(13.333), Inches(7.5)
TOTAL = 15


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- helpers
def runs(par, text, size, color=TEXT, bold=False):
    """Add text to a paragraph; **bold** segments are bold."""
    for i, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        r = par.add_run()
        r.text = part
        r.font.size, r.font.name = Pt(size), FONT
        r.font.bold = bold or i % 2 == 1
        r.font.color.rgb = rgb(color)


def box(slide, l, t, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, l, t, w, h)
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def text(slide, l, t, w, h, lines, size=16, color=TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=6, bullet=False):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, item in enumerate(lines if isinstance(lines, list) else [lines]):
        lvl = 0
        if isinstance(item, tuple):
            item, lvl = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        prefix = ("•  " if lvl == 0 else "–  ") if bullet else ""
        if lvl:
            p.level = 1
        runs(p, prefix + item, size - 2 * lvl, color if not lvl else MUTED)
    return tb


def picture(slide, name, l, t, max_w, max_h, caption=None, center=True):
    path = os.path.join(IMG, name)
    iw, ih = Image.open(path).size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    x = l + (max_w - w) // 2 if center else l
    y = t + (max_h - h) // 2 if center else t
    pic = slide.shapes.add_picture(path, x, y, w, h)
    pic.line.color.rgb = rgb(LINE); pic.line.width = Pt(0.75)
    if caption:
        text(slide, x, y + h + Inches(0.04), w, Inches(0.35), caption, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    return pic


def new_slide(prs, title, n, kicker=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    if kicker:
        text(s, Inches(0.6), Inches(0.28), Inches(9), Inches(0.35), kicker.upper(), size=12, color=MS_C)
    text(s, Inches(0.6), Inches(0.52), Inches(12.1), Inches(0.8), title, size=32, color=TEXT, space=0)
    text(s, Inches(12.0), Inches(7.05), Inches(1.0), Inches(0.3), f"{n} / {TOTAL}", size=11, color=MUTED, align=PP_ALIGN.RIGHT)
    return s


def takeaway(slide, msg):
    l, t, w, h = Inches(0.6), Inches(6.35), Inches(12.1), Inches(0.62)
    box(slide, l, t, w, h, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.15
    tb = text(slide, l + Inches(0.25), t, w - Inches(0.35), h, f"**Key takeaway:** {msg}", size=16, anchor=MSO_ANCHOR.MIDDLE, space=0)
    return tb


def notes(slide, summary, key, talk=()):
    """English notes, used only when a slide has no Hebrew notes in HE_NOTES."""
    tf = slide.notes_slide.notes_text_frame
    tf.text = f"SUMMARY: {summary}\n\nKEY TAKEAWAY: {key}"
    for line in talk:
        tf.add_paragraph().text = "- " + line


# Hebrew speaker notes, one entry per slide (what to say, ~1.5–2 minutes each).
HE_NOTES = [
    # 1
    """שלום לכולם. הפרויקט שלנו עוסק בסימון אוטומטי של גידולי מוח בסריקות MRI.
הרעיון המרכזי: משלבים שני מודלים. YOLO הוא מודל מהיר שמוצא את הגידול בכל פרוסה בנפרד, ו-MedSAM2 הוא מודל שמקבל את הניחוש של YOLO ומשפר אותו בעזרת הפרוסות השכנות, כלומר בתלת-ממד.
עבדנו על מאגר BraTS 2024 של גליומות במבוגרים: 1,621 סריקות.
בתמונה למטה רואים את כל התהליך על פרוסה אחת: משמאל התמונה שהמודל רואה, אחריה הרמז של YOLO, אחר כך המסכה של YOLO עצמו, ומימין התשובה של MedSAM2. ירוק = נכון, אדום = מיותר.""",
    # 2
    """למה בכלל צריך סימון אוטומטי?
גליומות הן גידולי המוח הממאירים הראשוניים הנפוצים ביותר במבוגרים.
רופאים מסמנים את הגידול ב-MRI כדי לתכנן ניתוח והקרנות וכדי לעקוב אחרי שינוי בין סריקות.
הבעיה: סריקה אחת היא יותר ממאה פרוסות, וכל אחת מסומנת ביד. זה איטי, וגם מומחים לא מסכימים ביניהם לגמרי.
המטרה: סימון תלת-ממדי מהיר ועקבי, שהרופא רק צריך לבדוק ולאשר.
מסר: סימון אוטומטי חוסך זמן למומחים ונותן מדידות עקביות.""",
    # 3
    """כאן רואים מטופל אחד, BraTS-GLI-00046-101, בפרוסה 115, בהגדלה על אזור הגידול. בפרוסה הזאת יש את כל ארבעת חלקי הגידול.
T1 בלי חומר ניגוד: הטבעת של הגידול הפעיל (המסומנת בסגול) כמעט לא בהירה. זו תמונת ה"לפני" שמשווים אליה. המודלים שלנו לא משתמשים בה.
T1C, אחרי הזרקת חומר ניגוד: אותה טבעת נדלקת בבהירות. זה הגידול הפעיל (ET). בפנים נשארת ליבה כהה, הרקמה המתה (NETC, באדום).
T2: בצקת ונוזלים בהירים. הקו הירוק מסמן את הבצקת (SNFH).
FLAIR: כמו T2, אבל הנוזל הרגיל של המוח (בחדרים) כהה, ולכן הבצקת בולטת יותר.
בצד ימין: הסימון של המומחה, כולל חלל הניתוח (RC) בצהוב.
אנחנו מסמנים את הגידול כולו, כלומר את כל החלקים יחד. המודלים שלנו מקבלים את T1C, T2 ו-FLAIR כתמונה צבעונית אחת.
המדד: 3D Dice, החפיפה בין המסכה של המודל לזו של המומחה בכל הנפח. 0 = אין חפיפה, 1 = מושלם.
מסר: אף רצף לבד לא מראה את כל הגידול, ולכן משלבים כמה רצפים.""",
    # 4
    """שני המודלים משלימים זה את זה.
YOLO11: מודל זיהוי מהיר שגם מצייר מסכה לכל צורה שהוא מוצא. הוא עובד על פרוסה אחת בכל פעם ולא יודע מה יש בפרוסות השכנות. לכל צורה הוא נותן ציון ביטחון בין 0 ל-1.
SAM2 של Meta יודע לסמן אובייקט לפי רמז: נקודה, מלבן או מסכה. הוא מתייחס לסריקה תלת-ממדית כמו לווידאו וזוכר את הפרוסות שכבר עבר. MedSAM2 הוא SAM2 שאומן בנוסף על תמונות רפואיות.
ל-MedSAM2 יש כ-39 מיליון משקלים, 70% מהם במקודד התמונה.
מסר: YOLO טוב במציאה, MedSAM2 טוב בתיקון בעזרת ההקשר התלת-ממדי.""",
    # 5
    """שאלת המחקר: האם MedSAM2, כשהוא מקבל רמז מ-YOLO בכל פרוסה, מסמן את הגידול טוב יותר מ-YOLO לבד?
כדי שההשוואה תהיה הוגנת שמרנו על כמה כללים:
המסכה של המומחה אף פעם לא מוצגת למודל בזמן חיזוי, היא רק "דף התשובות".
את ההגדרות בחרנו על קבוצות ולידציה, ואת קבוצת המבחן הרצנו פעם אחת בלבד.
החלוקה היא לפי אדם, לא לפי סריקה.
ומשווים מול הגרסה הכי טובה של YOLO, לא מול גרסה חלשה.
בנינו שלושה דברים: YOLO מאומן, MedSAM2 מאומן, ואפליקציית ווב לניתוח הנתונים והתוצאות.""",
    # 6
    """המאגר: 1,621 סריקות מ-731 אנשים. חילקנו לחמש קבוצות לפי אדם: אימון ו-ולידציה ל-YOLO, אימון ו-ולידציה ל-MedSAM2, וקבוצת מבחן של 324 סריקות.
הנקודה החשובה: אף אדם לא נמצא גם בקבוצה של YOLO וגם בקבוצה של MedSAM2. כך MedSAM2 מתאמן על רמזים ש-YOLO יצר לאנשים שהוא לא ראה, בדיוק כמו בשימוש אמיתי.
את הגדלים קבענו לפי עקומות למידה: MedSAM2 כמעט לא השתפר מעבר ל-500 סריקות, בעוד ש-YOLO עוד הרוויח מעוד נתונים, ולכן העברנו אליו 221 סריקות.
הסתייגות כנה: קבוצת המבחן חולקת חלק מהאנשים (בסריקות אחרות) עם קבוצות האימון, מהחלוקה המקורית. השארנו כך כדי שכל תוצאות המבחן יהיו ברות השוואה.""",
    # 7
    """כדי להכיר את הנתונים בנינו אפליקציית ווב עם ארבעה עמודים: ניתוח נתונים, תוצאות, YOLO ו-MedSAM2. כל גרף מסנן את כל השאר.
כמה עובדות: גודל גידול חציוני 59 מ"ל, 85% מהסריקות אחרי ניתוח, ול-73% יש כמה חלקים נפרדים. הגידולים מתחלקים שווה בערך בין צד שמאל (48%) לימין (47%).
מפת החום מראה איפה נמצאים מרכזי הגידולים, במבט מלמעלה, כאשר הצד הימני של המטופל מוצג משמאל, כמו בסריקה בבית חולים.
השתמשנו באפליקציה כדי לבדוק את החלוקה, למצוא מטופלים חריגים, ולראות איפה כל מודל נכשל.""",
    # 8
    """זה כל התהליך.
מתחילים מפרוסות ה-MRI כתמונה צבעונית. YOLO רץ על כל פרוסה ושומר כל צורה עם ביטחון של 0.05 ומעלה.
מהפלט שלו בונים "מפת רמז": ההסתברות לכל פיקסל, מומרת ל-logit ומוכפלת ב-8. בפרוסות ריקות ומחוץ לראש שמים "אין כאן גידול" ברור.
MedSAM2 מתחיל מהפרוסה שבה YOLO הכי בטוח, ואז עובר למעלה ולמטה. בכל פרוסה הוא רואה את התמונה, את הרמז, ואת הזיכרון: פרוסת ההתחלה ועוד 6 הפרוסות האחרונות.
בנוסף מריצים את כל המטופל גם בתמונת מראה (שמאל-ימין), הופכים בחזרה וממצעים.
בסוף: מעל 50% = גידול, ומוחקים חתיכות קטנות מ-100 ווקסלים. רק אז משווים למומחה.""",
    # 9
    """שלב 1: אימון YOLO.
לקחנו את yolo11m-seg, שאומן מראש על תמונות יומיומיות, ואימנו אותו ברזולוציה 512 על הפרוסות של 471 סריקות, כולל פרוסות ריקות, כדי שילמד גם מתי לא לצייר כלום.
את הסבב הטוב ביותר בחרנו על הוולידציה, והאימון נעצר מוקדם בסבב 32 מתוך 60.
בזמן חיזוי הוספנו שלושה דברים: ביטחון נמוך של 0.05, מיצוע עם תמונת מראה, והסרת חתיכות קטנות מ-200 ווקסלים. זה הוסיף כ-0.006.
התוצאה על המבחן: 3D Dice של 0.8834, ו-93.8% מפרוסות הגידול נמצאו.
הגרף משמאל מראה שביטחון נמוך יותר, כלומר לצייר יותר, נותן תוצאה טובה יותר.
מסר: זה בסיס חזק מאוד למודל דו-ממדי.""",
    # 10
    """שלב 2: איך הפלט של YOLO הופך לרמז ל-MedSAM2.
הרמז הוא ההסתברות של YOLO לכל פיקסל, מכל הצורות עם ביטחון 0.05 ומעלה. זה נותן ל-MedSAM2 יותר מידע, כולל ניחושים חלשים.
ממירים את ההסתברות ל-logit, הסולם ש-MedSAM2 עצמו משתמש בו, שבו 0 הוא 50%. הלוגיט מדגיש את הקצוות: הוא מבדיל בין "די בטוח" ל"בטוח מאוד". אחרי כפל ב-8 הטווח הוא בערך מינוס 44 עד 44.
בפרוסות ש-YOLO לא מצא בהן כלום ומחוץ לראש שמים מינוס 20: "אין כאן כלום".
בתמונה, משמאל לימין: התמונה, הרמז, המסכה של YOLO עצמו, והתשובה של MedSAM2. שימו לב ש-MedSAM2 תיקן חלק מהאזורים שבהם YOLO טעה.""",
    # 11
    """איך מאמנים את MedSAM2 לתקן ולא להעתיק.
התאמנו על 500 סריקות, בקטעים של 8 פרוסות, במשך 15 סבבים. אחרי כל סבב בדקנו על 80 סריקות ולידציה ושמרנו את הטוב ביותר.
אימנו רק את החלקים הקטנים: מפענח המסכה ומקודד הרמז בקצב 1e-4, והזיכרון בקצב איטי יותר, 1e-5. בסך הכול 11.7 מיליון משקלים, 30% מהמודל.
את מקודד התמונה הגדול השארנו קפוא. ניסינו לאמן גם אותו, וזה לא שיפר.
הטריק החשוב: ב-30% מהדוגמאות קלקלנו את הרמז בכוונה (הסרנו, הזזנו, הגדלנו, הקטנו), בזמן שהתשובה נשארת המסכה של המומחה. כך המודל לומד לתקן את YOLO ולא סתם להעתיק אותו.
בגרפים רואים שהציון מתייצב בערך מסבב 10.""",
    # 12
    """התוצאות על קבוצת המבחן, 324 מטופלים שאף מודל לא ראה.
MedSAM2 הגיע ל-0.8988, מול 0.8834 של הגרסה הכי טובה של YOLO.
השיפור הוא 0.0154, עם רווח סמך של 95% בין 0.012 ל-0.019. מבחן Wilcoxon נותן p קטן מאוד, כך שזה לא מקרי.
185 מטופלים השתפרו ביותר מ-0.01, ורק 18 הורעו.
MedSAM2 מצא 96.9% מפרוסות הגידול, והשאיר ריקות 96.4% מהפרוסות הריקות.
חשוב: את כל ההגדרות בחרנו על הוולידציה, ואת המבחן הרצנו פעם אחת.
מסר: השיפור עקבי וברור סטטיסטית, גם מול YOLO הכי טוב.""",
    # 13
    """איפה זה עובד ואיפה פחות.
בטבלה למעלה: הציון הממוצע לפי צד (שורות) וגודל גידול (עמודות). n הוא מספר המטופלים בכל משבצת.
גודל הגידול קובע הרבה יותר מהצד: בגידולים קטנים מ-10 מ"ל הציון בערך 0.67 עד 0.73, ובגידולים מעל 100 מ"ל בערך 0.93. שמאל וימין כמעט זהים.
דווקא בגידולים הקטנים MedSAM2 משפר הכי הרבה ביחס ל-YOLO.
הגרף למטה מראה מטופל אחד פרוסה אחר פרוסה: MedSAM2 עוקב אחרי הגידול בעזרת הזיכרון, ודוחה התראות שווא בודדות של YOLO.
הסתייגות: בחלק מהמשבצות יש מעט מאוד מטופלים, למשל בשורה "שני הצדדים" (17 בסך הכול), ולכן הן לא אמינות.""",
    # 14
    """מה ניסינו ומה למדנו.
התכנון הראשון שלנו נתן רמז רק לפרוסת ההתחלה, עם סבבי תיקון נוספים, והוא היה גרוע יותר מ-YOLO: 0.856 מול 0.864 על הוולידציה. כשעברנו לתת רמז בכל פרוסה, זה תוקן.
אחר כך בדקנו הגדרות: אילו חלקים לאמן, עוצמת הרמז, ערך "אין כאן כלום". כמעט כל שינוי הזיז את התוצאה ב-0.001 או פחות. רק "אין כאן כלום" חזק יותר (מינוס 20) היה שיפור אמיתי סטטיסטית, הירוק בגרף.
למה השיפורים קטנים? כי 96% מהטעויות שנשארו הן בשולי הגידול, ולא במציאת הגידול עצמו.
מסר: התכנון היה חשוב הרבה יותר מכיוונון ההגדרות, ומה שנשאר לשפר הוא דיוק הגבולות.""",
    # 15
    """עבודה עתידית, בשלושה כיוונים.
גבולות חדים יותר: למצע כמה מודלי MedSAM2 שכבר אימנו (אנסמבל), בלי אימון חדש; פונקציית הפסד שמתמקדת בגבולות; ורזולוציית פלט גבוהה יותר.
רמזים טובים יותר וגידולים קטנים: לשקלל כל פיקסל ברמז לפי הביטחון של הצורה; לתת ל-YOLO הקשר תלת-ממדי; ולהתמקד בגידולים מתחת ל-10 מ"ל.
מעבר לפרויקט: לסמן גם את חלקי הגידול (ET, NETC, SNFH) ולא רק את כולו; לבדוק על סריקות מבתי חולים אחרים; להסיר את החפיפה של אנשים בין המבחן לאימון; ולמדוד מהירות.
חשוב לזכור שגם מומחים לא מסכימים ביניהם לגמרי, ולכן השיפורים מכאן יהיו קטנים.
תודה רבה, אשמח לשאלות.""",
]


# ---------------------------------------------------------------- slides
def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # 1 ---- title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    text(s, Inches(0.7), Inches(1.0), Inches(8), Inches(0.4), "CLASS PROJECT · OCTOBER 2026", size=13, color=MS_C)
    text(s, Inches(0.7), Inches(1.45), Inches(11.5), Inches(1.6), "Brain Tumour Segmentation on MRI with YOLO + MedSAM2", size=40, space=0)
    text(s, Inches(0.7), Inches(3.2), Inches(11.5), Inches(0.9),
         "A fast 2D detector finds the tumour; a promptable video model refines it in 3D. BraTS 2024 adult glioma, 1,621 MRI scans.",
         size=18, color=MUTED)
    picture(s, "ms_panels.png", Inches(0.7), Inches(4.25), Inches(11.9), Inches(2.4),
            caption="One slice: the MRI picture · YOLO's hint · YOLO's own mask · MedSAM2's answer (green = found, red = extra)")
    text(s, Inches(0.7), Inches(6.95), Inches(8), Inches(0.35), "Presented by: ______________", size=13, color=MUTED)
    notes(s, "Introduces the project: segmenting brain tumours on MRI by combining YOLO and MedSAM2.",
          "We combine a fast 2D detector with a 3D-aware promptable model to outline brain tumours.",
          ["The picture shows the whole pipeline on one slice, left to right."])

    # 2 ---- background: why
    s = new_slide(prs, "Why segment brain tumours automatically?", 2, "Background")
    text(s, Inches(0.6), Inches(1.6), Inches(7.2), Inches(4.6), [
        "**Gliomas** are the most common malignant primary brain tumours in adults.",
        "Doctors outline the tumour on MRI to **plan surgery and radiotherapy** and to **track change** between scans.",
        "Outlining by hand is **slow**: a scan has well over a hundred slices, each drawn separately.",
        "Experts **don't agree perfectly** with each other, so outlines vary between readers.",
        "Automation goal: a **fast, consistent 3D outline** a doctor only needs to check.",
    ], size=18, bullet=True, space=12)
    picture(s, "mri_flair_seg.png", Inches(8.3), Inches(1.55), Inches(4.4), Inches(4.1),
            caption="An expert outline on FLAIR (green = swelling, yellow = surgery cavity)")
    takeaway(s, "Automatic 3D segmentation saves expert time and makes tumour measurements consistent.")
    notes(s, "Explains the clinical need for automatic tumour outlining.",
          "Automatic 3D segmentation saves expert time and makes tumour measurements consistent.",
          ["Manual contouring is the bottleneck; variability between experts limits how 'perfect' any model can look."])

    # 3 ---- background: MRI sequences, one patient
    s = new_slide(prs, "Each MRI sequence shows a different part of the tumour", 3,
                  "Background · patient BraTS-GLI-00046-101, slice 115, zoomed on the tumour")
    panels = [("mod_t1.png", "T1 (no contrast)", "Rim (ET, violet) barely visible. Not used by our models."),
              ("mod_t1c.png", "T1C (with contrast)", "Rim lights up (ET, violet); dead core dark (NETC, red)."),
              ("mod_t2.png", "T2", "Swelling and fluid bright (SNFH, green)."),
              ("mod_flair.png", "FLAIR", "Like T2, but brain fluid is dark: swelling stands out."),
              ("mod_labels.png", "Expert labels", "NETC red · SNFH green · ET violet · RC (cavity) yellow")]
    pw, step, x0, top = Inches(2.05), Inches(2.32), Inches(0.985), Inches(1.45)
    for i, (f, head, cap) in enumerate(panels):
        l = x0 + i * step
        pic = picture(s, f, l, top, pw, Inches(3.2), center=False)
        y = pic.top + pic.height + Inches(0.06)
        text(s, l, y, pw, Inches(0.3), f"**{head}**", size=13, align=PP_ALIGN.CENTER, space=0)
        text(s, l, y + Inches(0.28), pw, Inches(0.6), cap, size=11, color=MUTED, align=PP_ALIGN.CENTER, space=0)
    text(s, Inches(0.6), Inches(5.45), Inches(12.1), Inches(0.8), [
        "Target: the **whole tumour** (all four parts). Our models see T1C, T2 and FLAIR as one colour picture.",
        "Score: **3D Dice** = 2 × overlap ÷ (model + expert), from 0 to 1, per patient.",
    ], size=14, bullet=True, space=3)
    takeaway(s, "No single sequence shows the whole tumour, so the models look at several sequences together.")
    notes(s, "Shows one patient in four MRI sequences and which tumour part each sequence reveals.",
          "No single sequence shows the whole tumour, so the models look at several sequences together.")

    # 4 ---- background: models
    s = new_slide(prs, "Two models with different strengths", 4, "Background")
    for i, (name, col, items) in enumerate([
        ("YOLO11-seg  (Ultralytics)", YOLO_C, [
            "Fast object detector that also draws a mask for each shape it finds.",
            "Works on **one 2D slice at a time**, no knowledge of neighbouring slices.",
            "Gives a **confidence** for every shape (0–1).",
            "We use the medium model (yolo11m-seg), pretrained on everyday photos."]),
        ("SAM2 → MedSAM2", MS_C, [
            "**Segment Anything 2** (Meta): outlines an object from a **prompt** (point, box or mask).",
            "Treats a 3D scan **like a video**: remembers the slices it already did.",
            "**MedSAM2** = SAM2 further trained on medical images and scans.",
            "39 M weights: image encoder 70%, memory 19%, mask decoder 11%, prompt encoder <0.1%."])]):
        l = Inches(0.6 + i * 6.15)
        box(s, l, Inches(1.6), Inches(5.95), Inches(4.5), fill=PANEL)
        text(s, l + Inches(0.25), Inches(1.8), Inches(5.5), Inches(0.5), f"**{name}**", size=20, color=col)
        text(s, l + Inches(0.25), Inches(2.45), Inches(5.5), Inches(3.6), items, size=16, bullet=True, space=10)
    takeaway(s, "YOLO is good at finding the tumour on a slice; MedSAM2 can refine it using the slices around it.")
    notes(s, "Introduces the two models and why they complement each other.",
          "YOLO is good at finding the tumour on a slice; MedSAM2 can refine it using the slices around it.",
          ["Weight counts measured on the actual checkpoint: 38,962,498 in total."])

    # 5 ---- question
    s = new_slide(prs, "Research question and ground rules", 5, "The project")
    box(s, Inches(0.6), Inches(1.6), Inches(12.1), Inches(1.1), fill=PANEL)
    text(s, Inches(0.85), Inches(1.6), Inches(11.6), Inches(1.1),
         "Can MedSAM2, prompted by YOLO on every slice, outline the whole tumour **better than YOLO alone** (3D Dice)?",
         size=21, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(0.6), Inches(3.0), Inches(6.0), Inches(0.45), "**Ground rules**", size=19, color=MS_C)
    text(s, Inches(0.6), Inches(3.5), Inches(6.0), Inches(2.7), [
        "The expert mask is never shown to a model when it predicts.",
        "Settings are chosen on validation pools; test is scored once.",
        "Patients are split by person, never by scan.",
        "We compare against **YOLO's own best setup**, not a weak baseline.",
    ], size=16, bullet=True, space=8)
    text(s, Inches(6.9), Inches(3.0), Inches(5.8), Inches(0.45), "**What we built**", size=19, color=YOLO_C)
    text(s, Inches(6.9), Inches(3.5), Inches(5.8), Inches(2.7), [
        "A fine-tuned YOLO (stage 1).",
        "A fine-tuned MedSAM2 that reads YOLO's hints (stage 2).",
        "An interactive web app to explore the data and every result.",
    ], size=16, bullet=True, space=8)
    takeaway(s, "A fair two-stage test: does adding MedSAM2 beat the best YOLO we could build?")
    notes(s, "States the research question and the rules that keep the comparison fair.",
          "A fair two-stage test: does adding MedSAM2 beat the best YOLO we could build?")

    # 6 ---- data split
    s = new_slide(prs, "Data: 1,621 scans from 731 people, split by person", 6, "Data")
    picture(s, "an_manifest.png", Inches(0.6), Inches(1.55), Inches(5.6), Inches(4.6))
    text(s, Inches(6.5), Inches(1.6), Inches(6.2), Inches(4.6), [
        "Five pools: YOLO train 471 · YOLO val 81 · MedSAM2 train 500 · MedSAM2 val 245 · **test 324** scans.",
        "**No person is in both a YOLO pool and a MedSAM2 pool**, so MedSAM2 trains on hints YOLO made for people it never saw, like real use.",
        "Pool sizes set by learning curves: MedSAM2 barely improved past 500 patients (250 → 500 → 721: 0.884 → 0.886 → 0.886), while YOLO still gained from more data.",
        "Caveat: test shares some people (other scans) with training pools, from the original split; kept so all test scores stay comparable.",
    ], size=16, bullet=True, space=10)
    takeaway(s, "Keeping the two models' people apart makes MedSAM2's training hints realistic.")
    notes(s, "Shows how the data was split into five pools and why.",
          "Keeping the two models' people apart makes MedSAM2's training hints realistic.",
          ["221 scans were moved from MedSAM2 train to YOLO train on 2026-09-29 after the learning curves."])

    # 7 ---- analytics app
    s = new_slide(prs, "Knowing the data: an interactive analytics app", 7, "Data")
    picture(s, "an_kpis.png", Inches(0.6), Inches(1.5), Inches(12.1), Inches(1.9))
    picture(s, "an_heat.png", Inches(0.6), Inches(3.5), Inches(4.0), Inches(2.75))
    text(s, Inches(4.9), Inches(3.55), Inches(7.8), Inches(2.7), [
        "4 pages: **Analytics**, **Results**, **YOLO**, **MedSAM2**; every chart filters the others.",
        "Median tumour 59 mL; **85%** of scans are after surgery; **73%** have several separate pieces.",
        "Tumours split evenly between the left (48%) and right (47%) brain.",
        "Used to check the split, find unusual patients and see where each model fails.",
    ], size=16, bullet=True, space=9)
    takeaway(s, "Understanding sizes, sides and pieces explained where the models struggle.")
    notes(s, "Presents the web app built to explore the dataset and the results.",
          "Understanding sizes, sides and pieces explained where the models struggle.",
          ["Heat map: where tumour centres are, seen from above (patient's right on the viewer's left)."])

    # 8 ---- pipeline diagram
    s = new_slide(prs, "The pipeline: YOLO finds, MedSAM2 refines", 8, "Method")
    steps = [("MRI slices", "T1C · T2 · FLAIR\nas one picture", PANEL, TEXT),
             ("YOLO11-seg", "every slice\nconfidence ≥ 0.05", YOLO_C, "FFFFFF"),
             ("Hint map", "logit × 8\nempty = 'no'", PANEL, TEXT),
             ("MedSAM2", "up & down from start\nmemory: start + 6", MS_C, "FFFFFF"),
             ("Mirror check", "mirror patient\nand average", PANEL, TEXT),
             ("Final 3D mask", "> 50% = tumour\nno pieces < 100 vox", PANEL, TEXT)]
    bw, bh, gap, top = Inches(1.78), Inches(1.9), Inches(0.29), Inches(2.0)
    for i, (head, sub, fill, ink) in enumerate(steps):
        l = Inches(0.6) + i * (bw + gap)
        b = box(s, l, top, bw, bh, fill=fill, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        b.adjustments[0] = 0.08
        text(s, l + Inches(0.08), top + Inches(0.15), bw - Inches(0.16), Inches(0.45), f"**{head}**", size=16, color=ink, align=PP_ALIGN.CENTER)
        text(s, l + Inches(0.08), top + Inches(0.7), bw - Inches(0.16), Inches(1.1), sub, size=13, color=ink, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            a = box(s, l + bw + Inches(0.03), top + bh / 2 - Inches(0.14), gap - Inches(0.06), Inches(0.28), fill=MUTED, shape=MSO_SHAPE.RIGHT_ARROW)
    text(s, Inches(0.6), Inches(4.25), Inches(12.1), Inches(2.0), [
        "**Stage 1 (YOLO)** runs once per patient; its hint maps are saved.",
        "**Stage 2 (MedSAM2)** gets the picture **and** YOLO's hint on every slice, plus its memory of the slices it just did.",
        "Only at the very end is the 3D mask compared with the expert's (3D Dice).",
    ], size=17, bullet=True, space=10)
    takeaway(s, "MedSAM2 doesn't search from scratch; it corrects YOLO's guess using the neighbouring slices.")
    notes(s, "Walks through the full inference pipeline from MRI slices to the final 3D mask.",
          "MedSAM2 doesn't search from scratch; it corrects YOLO's guess using the neighbouring slices.",
          ["Start slice = the slice where YOLO is most confident (area × score).",
           "Mirror check: the mirrored copy runs in the same batch, so it costs almost nothing."])

    # 9 ---- YOLO
    s = new_slide(prs, "Stage 1: fine-tuning YOLO", 9, "Method · Results")
    picture(s, "yo_eval_top.png", Inches(0.6), Inches(1.45), Inches(12.1), Inches(3.35))
    text(s, Inches(0.6), Inches(4.95), Inches(6.0), Inches(1.35), [
        "yolo11m-seg at 512 px, trained on the slices of 471 scans, **including empty ones**.",
        "Best round picked on YOLO val; stopped early at round 32 of 60.",
    ], size=15, bullet=True, space=6)
    text(s, Inches(6.8), Inches(4.95), Inches(5.9), Inches(1.35), [
        "Extras: confidence 0.05, mirror averaging, drop pieces < 200 voxels (+0.006 val).",
        "**Test 3D Dice 0.8834** (median 0.909); found 93.8% of tumour slices.",
    ], size=15, bullet=True, space=6)
    takeaway(s, "A strong 2D baseline: 0.883 3D Dice on 324 unseen test patients.")
    notes(s, "Describes how YOLO was fine-tuned and how well it does on its own.",
          "A strong 2D baseline: 0.883 3D Dice on 324 unseen test patients.",
          ["The confidence chart shows lower confidence = more drawn = better 3D Dice, up to 0.05–0.1."])

    # 10 ---- hint
    s = new_slide(prs, "Stage 2: turning YOLO's output into a MedSAM2 hint", 10, "Method")
    picture(s, "ms_viewer.png", Inches(0.6), Inches(1.5), Inches(12.1), Inches(3.0))
    text(s, Inches(0.6), Inches(4.6), Inches(6.0), Inches(1.7), [
        "Hint = YOLO's **probability per pixel** from every shape at confidence ≥ 0.05.",
        "Turned into a **logit × 8** (−44 … +44): MedSAM2's own scale, where 0 = 50%.",
    ], size=16, bullet=True, space=8)
    text(s, Inches(6.8), Inches(4.6), Inches(5.9), Inches(1.7), [
        "Empty slices and outside the head get a firm **'nothing here'** (−20).",
        "Every slice is hinted; memory carries the start slice + the last 6 slices.",
    ], size=16, bullet=True, space=8)
    takeaway(s, "YOLO's soft guess plus 3D memory lets MedSAM2 fix single-slice mistakes.")
    notes(s, "Explains how YOLO's output becomes the prompt MedSAM2 reads on every slice.",
          "YOLO's soft guess plus 3D memory lets MedSAM2 fix single-slice mistakes.",
          ["Picture (left to right): input picture, YOLO's hint, YOLO's own mask, MedSAM2's answer.",
           "YOLO's own mask uses only shapes ≥ 0.25 and pixels > 50%; it is only for comparison."])

    # 11 ---- training
    s = new_slide(prs, "Training MedSAM2 to fix, not copy", 11, "Method")
    picture(s, "ms_training_curves.png", Inches(0.6), Inches(1.5), Inches(7.0), Inches(4.7))
    text(s, Inches(7.9), Inches(1.6), Inches(4.8), Inches(4.7), [
        "500 training scans, clips of 8 slices, 15 rounds; best round checked on 80 val scans.",
        "**Trained:** mask decoder + prompt encoder (lr 1e-4) and memory (1e-5) = 11.7 M weights (30%).",
        "**Frozen:** the image encoder (70% of weights); training it too gave no gain.",
        "**30% of training hints damaged on purpose** (dropped, shifted, grown, shrunk) while the answer stays the expert's.",
        "Loss: Dice + cross-entropy (+ small IoU and object terms).",
    ], size=15, bullet=True, space=9)
    takeaway(s, "Train only the small parts that read the hint and carry memory; keep the big image encoder frozen.")
    notes(s, "Explains what parts of MedSAM2 are trained and how training teaches correction.",
          "Train only the small parts that read the hint and carry memory; keep the big image encoder frozen.",
          ["Check score levels off around round 10 at about 0.90."])

    # 12 ---- results
    s = new_slide(prs, "Results: MedSAM2 beats YOLO's best on test", 12, "Results")
    picture(s, "chart_final_comparison.png", Inches(0.6), Inches(1.5), Inches(6.6), Inches(4.7))
    text(s, Inches(7.5), Inches(1.6), Inches(5.2), Inches(4.7), [
        "**0.8988** vs **0.8834** mean 3D Dice on the same 324 test patients.",
        "Gain **+0.0154** (95% CI +0.012 … +0.019); Wilcoxon p < 10⁻³³.",
        "**185 patients better, 18 worse** (by more than 0.01); the rest about the same.",
        "Found 96.9% of tumour slices and kept 96.4% of empty slices empty.",
        "Settings were picked on val; test was scored once.",
    ], size=16, bullet=True, space=10)
    takeaway(s, "Adding MedSAM2 gives a consistent, statistically clear gain over the best YOLO.")
    notes(s, "Presents the final test results against YOLO's best setup.",
          "Adding MedSAM2 gives a consistent, statistically clear gain over the best YOLO.",
          ["The fair comparison is against 'YOLO best' (0.8834), not 'YOLO alone' (0.8828)."])

    # 13 ---- where
    s = new_slide(prs, "Where it works, and where it struggles", 13, "Results")
    picture(s, "chart_side_size.png", Inches(0.6), Inches(1.45), Inches(6.9), Inches(2.95),
            caption="Mean MedSAM2 3D Dice by side (rows) and tumour size (columns), test")
    picture(s, "ms_patient_chart.png", Inches(0.6), Inches(4.8), Inches(6.9), Inches(1.45))
    text(s, Inches(7.8), Inches(1.6), Inches(4.9), Inches(4.6), [
        "**Size matters, side doesn't:** about 0.67–0.73 for tumours under 10 mL vs about 0.93 for 100 mL and up; left ≈ right.",
        "The **biggest gains** over YOLO are on **small tumours**.",
        "Slice by slice (bottom): MedSAM2 follows the tumour through memory and **rejects YOLO's isolated false alarms**.",
        "Few patients in some groups (e.g. 'both sides': 17), so those squares are noisy.",
    ], size=15, bullet=True, space=10)
    takeaway(s, "Tiny tumours remain the weak spot, and they are where MedSAM2 helps most.")
    notes(s, "Breaks results down by tumour size and side and shows one patient slice by slice.",
          "Tiny tumours remain the weak spot, and they are where MedSAM2 helps most.")

    # 14 ---- lessons
    s = new_slide(prs, "What we tried and what we learned", 14, "Experiments")
    picture(s, "chart_experiments.png", Inches(0.6), Inches(1.5), Inches(6.2), Inches(3.3),
            caption="Settings experiments (val): only the green change was statistically real")
    picture(s, "chart_errors.png", Inches(7.0), Inches(1.5), Inches(5.7), Inches(3.3),
            caption="Where the remaining wrong voxels are (val)")
    text(s, Inches(0.6), Inches(5.15), Inches(12.1), Inches(1.15), [
        "First design (only the start slice prompted, plus extra correction rounds) made things **worse** than YOLO (0.856 vs 0.864, val); hinting **every** slice fixed it.",
        "Most settings changes moved Dice by ≤ 0.001: **96% of the remaining error is at the tumour edge**, not in finding the tumour.",
    ], size=15, bullet=True, space=6)
    takeaway(s, "Design mattered more than tuning; what's left to win is edge precision.")
    notes(s, "Summarises the experiments: the redesign that worked, tuning results, and where errors remain.",
          "Design mattered more than tuning; what's left to win is edge precision.",
          ["The old-design numbers were on the earlier pools, so they are compared within that setup only."])

    # 15 ---- future work
    s = new_slide(prs, "Future work", 15, "Next steps")
    cols = [("Sharper edges", MS_C, ["Average several trained MedSAM2 runs (an ensemble), with no new training.",
                                     "Edge-focused training loss.",
                                     "Finer output: scans are 182 × 218, the model works at 512 × 512."]),
            ("Hints & small tumours", YOLO_C, ["Weigh hint pixels by shape confidence ('max_weighted').",
                                               "Give YOLO 3D context (neighbouring slices).",
                                               "Train more on tumours under 10 mL."]),
            ("Beyond this project", GOOD, ["Predict the tumour parts (ET, NETC, SNFH), not only the whole tumour.",
                                           "Test on other hospitals' scans.",
                                           "Remove the test/train people overlap; measure speed."])]
    for i, (head, col, items) in enumerate(cols):
        l = Inches(0.6 + i * 4.1)
        box(s, l, Inches(1.6), Inches(3.9), Inches(4.4), fill=PANEL)
        text(s, l + Inches(0.2), Inches(1.8), Inches(3.5), Inches(0.5), f"**{head}**", size=19, color=col)
        text(s, l + Inches(0.2), Inches(2.45), Inches(3.5), Inches(3.5), items, size=16, bullet=True, space=12)
    takeaway(s, "Next gains come from sharper edges and small tumours, then proving it on outside data.")
    notes(s, "Lists the most promising next steps and the wider follow-ups.",
          "Next gains come from sharper edges and small tumours, then proving it on outside data.",
          ["Experts themselves don't agree perfectly, so expect small gains from here."])

    for slide, he in zip(prs.slides, HE_NOTES):
        slide.notes_slide.notes_text_frame.text = he
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
